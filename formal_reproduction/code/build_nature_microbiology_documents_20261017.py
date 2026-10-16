"""Write the replacement manuscript and submission documents from frozen tables."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt


OUT = Path(r"D:\ProtoEss_NatureMicrobiology_submission_20261017")
FIG = OUT / "figures"
AUTHORS = "Yuan-Nong Ye, Lu-Yao Tian, Ying Qiu, Lu Lei, Peng Guo and Zhu Zeng"
AFFILIATIONS = (
    "Guizhou Key Laboratory of Microbiome and Infectious Disease Prevention and Control, "
    "School of Biology and Engineering, Guizhou Medical University, Guiyang, China; "
    "Laboratory of Intelligent Analysis and Mining for Biomedical Big Data, Department of Medical Artificial Intelligence, "
    "School of Biology and Engineering (Modern Industrial College of Health Medicine), Guizhou Medical University, Guiyang, China."
)


def read_metrics() -> pd.DataFrame:
    return pd.read_csv(OUT / "metrics" / "frozen_fivefold_nine_route_metrics.tsv", sep="\t")


def write_text(path: Path, text: str) -> None:
    path.write_text(text.strip() + "\n", encoding="utf-8")


def docx_from_markdownish(path: Path, title: str, markdown: str, figures: list[tuple[str, str]]) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(AUTHORS).bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(AFFILIATIONS)
    paragraph_lines: list[str] = []
    def flush_paragraph() -> None:
        if paragraph_lines:
            text = " ".join(line.strip() for line in paragraph_lines).strip()
            if text:
                doc.add_paragraph(text)
            paragraph_lines.clear()
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
        elif stripped.startswith("### "):
            flush_paragraph()
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            flush_paragraph()
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            flush_paragraph()
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("- "):
            flush_paragraph()
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped.startswith("Table "):
            flush_paragraph()
            p = doc.add_paragraph()
            p.add_run(stripped).bold = True
        else:
            paragraph_lines.append(stripped)
    flush_paragraph()
    for filename, caption in figures:
        path_img = FIG / filename
        if path_img.exists():
            doc.add_picture(str(path_img), width=Inches(6.5))
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
    doc.save(path)


def main() -> None:
    for sub in ["manuscript", "supplementary_information", "submission"]:
        (OUT / sub).mkdir(parents=True, exist_ok=True)
    metrics = read_metrics()
    tests = metrics[metrics["split"] == "test"].set_index("species")
    loso = pd.read_csv(OUT / "loso" / "four_species_clade_loso_metrics.tsv", sep="\t")
    app = pd.read_csv(OUT / "predictions" / "standardized_genomewide_predictions.tsv", sep="\t")
    counts = app.assign(e=app.binary_call.eq("essential")).groupby("species").e.sum().to_dict()
    freeze = json.loads((OUT / "manifest" / "final_freeze_manifest.json").read_text(encoding="utf-8"))

    abstract = (
        "Genetic essentiality is measured under species- and condition-specific experimental constraints, "
        "which complicates direct transfer of parasite vulnerability maps. We built a frozen five-fold, nine-route "
        "LightGBM framework using the original 95%-redundancy-reduced, label-stratified random split and 628 raw features. "
        f"Held-out AUROC/AUPRC values were {tests.loc['Plasmodium_falciparum','AUROC']:.3f}/{tests.loc['Plasmodium_falciparum','AUPRC']:.3f} for P. falciparum, "
        f"{tests.loc['Plasmodium_berghei','AUROC']:.3f}/{tests.loc['Plasmodium_berghei','AUPRC']:.3f} for P. berghei and "
        f"{tests.loc['Plasmodium_knowlesi','AUROC']:.3f}/{tests.loc['Plasmodium_knowlesi','AUPRC']:.3f} for P. knowlesi. "
        "A four-species leave-one-species-out analysis, including Toxoplasma gondii, showed recipient-dependent transfer. "
        "Two independent P. knowlesi screens conducted under different red-cell conditions showed 12.57% label discordance "
        "among definite matched calls and a positive continuous-score association. We therefore interpret discordant calls as "
        "species- and condition-associated hypotheses rather than immutable species properties. The release provides a frozen, "
        "hash-indexed atlas, reproducible code and explicit limitations for computational prioritisation."
    )

    body = f"""
# Abstract
{abstract}

## Introduction
Genome-scale perturbation screens have transformed parasite biology, but essentiality is not an invariant gene attribute. It depends on genetic background, developmental state, host cell, assay design and the perturbation readout. These dependencies matter when experimental resources from different parasites are combined to infer missing labels or prioritize antimalarial targets.

Here we separate three questions that are often conflated: how well a model ranks labels within a species with a compatible screen; how much ranking transfers when the target species is completely withheld; and how much apparent disagreement can arise when the same species is screened under different conditions. We focus on P. falciparum, P. berghei and P. knowlesi, and include T. gondii in the leave-one-species-out analysis as a distant apicomplexan control.

## Results
### A frozen nine-route model ranks held-out malaria essentiality labels
The formal release uses the original 95%-redundancy-reduced, label-stratified random train, validation and test allocation. The model consumes 628 raw features without PCA or manual removal of homology30/50. Nine route probabilities are combined by the fixed LightGBM meta-model. Fold-local dynamic features, including label-derived homology, MMseqs and PPI summaries, are reconstructed from the fitting partition only.

The frozen test AUROC/AUPRC values were {tests.loc['Plasmodium_falciparum','AUROC']:.3f}/{tests.loc['Plasmodium_falciparum','AUPRC']:.3f} for Pf, {tests.loc['Plasmodium_berghei','AUROC']:.3f}/{tests.loc['Plasmodium_berghei','AUPRC']:.3f} for Pb and {tests.loc['Plasmodium_knowlesi','AUROC']:.3f}/{tests.loc['Plasmodium_knowlesi','AUPRC']:.3f} for Pk (Fig. 2). Thresholds were fixed from training-only probability analyses and are recorded with every score. The test labels were not used for current fitting, thresholding or candidate selection.

These metrics are a final evaluation of a frozen route, not a fully untouched prospective estimate. The historical route profile was chosen after prior inspection of Pk results, and the meta-level cross-validation is not fully nested. We therefore report the historical exposure explicitly and avoid calling the scores blind validation.

### Leave-one-species-out transfer includes Tg
In the four-species clade analysis, the target species was removed from training and the remaining species were used to score it. Test AUROC/AUPRC values were {loso.loc[loso.held_out=='Plasmodium_falciparum','test_AUROC'].iloc[0]:.3f}/{loso.loc[loso.held_out=='Plasmodium_falciparum','test_AUPRC'].iloc[0]:.3f} for Pf, {loso.loc[loso.held_out=='Plasmodium_berghei','test_AUROC'].iloc[0]:.3f}/{loso.loc[loso.held_out=='Plasmodium_berghei','test_AUPRC'].iloc[0]:.3f} for Pb, {loso.loc[loso.held_out=='Plasmodium_knowlesi','test_AUROC'].iloc[0]:.3f}/{loso.loc[loso.held_out=='Plasmodium_knowlesi','test_AUPRC'].iloc[0]:.3f} for Pk and {loso.loc[loso.held_out=='Toxoplasma_gondii','test_AUROC'].iloc[0]:.3f}/{loso.loc[loso.held_out=='Toxoplasma_gondii','test_AUPRC'].iloc[0]:.3f} for Tg (Fig. 3). Tg is therefore included as a complete fourth held-out species, not omitted from the transfer analysis.

LOSO scores answer a different question from matched-species atlas performance. They quantify ranking transfer under target-label withholding and should not be used as the species-matched thresholds for the three-species atlas.

### Independent Pk screens quantify condition-associated disagreement
The two P. knowlesi screen resources were aligned through one-to-one orthology. The primary definite-call comparison contained 3,501 matched genes and 440 label flips (12.57%; 95% CI 11.51–13.71%). Continuous scores from the two screens showed Spearman rho 0.744 (95% bootstrap CI 0.730–0.757). The paired comparison with exact Pf–Pb–Pk triplets showed Pf–Pk observed-label discordance of 19.65% versus 13.10% within the Pk dual-screen comparison, whereas Pb–Pk discordance was 6.77% versus 13.10%. These are descriptive matched-group contrasts, not causal decompositions of evolution and assay noise.

The present release does not claim exact A1H sequence-level prediction: the audited directories did not contain a complete A1H sequence-level feature matrix. The Pk human-red-cell module is therefore reported as a retrospective external-condition, ortholog-mapped analysis. The Oberstaller labels are not used as an independent evaluation target for a model trained with those same labels.

### Genome-scale application and comparative hypotheses
The frozen matched-species application yielded {counts.get('Plasmodium_falciparum', 0):,}, {counts.get('Plasmodium_berghei', 0):,} and {counts.get('Plasmodium_knowlesi', 0):,} predicted-essential calls for Pf, Pb and Pk, respectively (Fig. 4). Each record retains its score, threshold, model and feature version, screen-condition metadata and evidence class.

Across exact Pf–Pb–Pk orthogroups, the previously defined 535 directional groups and 46 refit-stable groups are retained as post-score comparative hypotheses. We use “species- and condition-associated essentiality differences” rather than “lineage-specific essentiality” because the dual-screen analysis demonstrates that experimental condition can change calls within a species. Homology-threshold sensitivity analyses remain supplementary and do not change the main random split.

Drug-target and host-risk overlays are applied after score generation. They are computational prioritisation layers and do not establish experimental essentiality, binding, selectivity, safety or efficacy.

## Discussion
The main result is methodological and biological: cross-species essentiality inference is useful but bounded. The frozen model integrates sequence, embedding, annotation, homology and network information, yet its transfer is asymmetric across recipient species. The inclusion of Tg makes that asymmetry explicit rather than treating only the three malaria parasites as a closed universe.

The Pk dual-screen analysis changes the interpretation of directional groups. A disagreement between species labels may reflect evolutionary background, host-cell condition, parasite strain, assay workflow or combinations of these factors. The atlas therefore preserves screen condition and evidence provenance and presents the 535 and 46 groups as testable hypotheses. It does not treat a computational binary call as a replacement for perturbation experiments.

The principal limitations are historical exposure of the selected route to Pk results, incomplete nesting of meta-level cross-validation, heterogeneous source screens and the lack of exact A1H feature-level reconstruction. These limitations are encoded in the freeze manifest and in the public model card. Future work should prospectively freeze the route before inspecting any target-species labels, obtain complete A1H sequence features, and test high-margin candidates in matched-stage perturbation experiments.

## Methods
### Data partition and feature construction
The main analysis reuses the original 95%-deduplicated, label-stratified random train/validation/test allocation. The 628-dimensional matrix is used at raw scale. No PCA, post-hoc feature deletion or test-set-driven feature selection is performed in the release run. Dynamic label-derived features are rebuilt within fitting folds. Matrix, identifier and prediction hashes are recorded in `manifest/final_freeze_manifest.json`.

### Five-fold nine-route LightGBM
Five-fold OOF probabilities are generated with shuffle=True and random_state=42. The nine route probabilities are passed to species-specific fixed LightGBM meta-model candidates selected from the prior training-CV profile table. Thresholds are derived from training-only probability analyses. Validation and test labels are used only for the corresponding final metric calculation.

### Leave-one-species-out analysis
Each of Pf, Pb, Pk and Tg is removed in turn. The remaining clade data are used for model fitting and the held-out species is scored without target labels. The reported LOSO table is a companion analysis and retains the historical clade full-stack model version in its metadata.

### Pk dual-screen analysis
Pk gene calls from the two screen conditions are aligned by the audited one-to-one orthology map. Flip rates use definite calls as the primary analysis; a tentative-call sensitivity analysis is also provided. Continuous association is evaluated by Spearman correlation with bootstrap confidence intervals. Since the exact A1H feature matrix was unavailable, the analysis is not described as exact human-RBC sequence-level prediction.

### Computational prioritisation
Genome-wide scores, binary calls and thresholds are exported using the stable schema `species, gene_id, score, binary_call, threshold, model_version, feature_version, screen_condition, evidence_class`. Drug-target, host-context and structural annotations are applied after model scoring.

## Data availability
All frozen tables, source data, manifests, figures and reproducibility scripts are versioned in the release bundle associated with `v20261017_5fold_9route_lightgbm`. The public website retains the legacy v48.1 layer and exposes the new version as a separate release. The final manifest contains SHA-256 hashes for published files.

## Code availability
The reproducibility repository contains the packaging, figure-generation and audit scripts. Public release code is intended to regenerate the key metrics and figures from the frozen source artifacts. The historical host-risk atlas repository is not modified by this release.

## Author contributions
Y.-N.Y. and Z.Z. conceived and supervised the study. L.-Y.T., Y.Q., L.L. and P.G. contributed to data curation, model analysis, evaluation and resource construction. All authors reviewed and approved the manuscript.

## Competing interests
The authors declare no competing interests.

## Funding
No applicable funding was declared in the source submission; this statement should be updated if institutional or grant support applies.

## Reporting summary
The study uses publicly available parasite genetic-screen labels, species-specific protein and annotation features, a fixed random split after 95% redundancy reduction, five-fold OOF training and a separate four-species LOSO analysis. No new animal or human experiments were performed. The Pk dual-screen comparison is retrospective and condition-associated.
"""

    write_text(OUT / "manuscript" / "Nature_Microbiology_replacement_manuscript.md", f"# Species- and condition-aware essentiality mapping across malaria parasites\n\n{AUTHORS}\n\n{body}")
    figs = [
        ("Figure_1_workflow.png", "Fig. 1 | Frozen analysis architecture and evidence boundaries."),
        ("Figure_2_main_performance.png", "Fig. 2 | Five-fold nine-route LightGBM held-out performance."),
        ("Figure_3_loso_four_species.png", "Fig. 3 | Four-species leave-one-species-out transfer, including Tg."),
        ("Figure_4_genomewide_application.png", "Fig. 4 | Genome-wide matched-species computational application."),
        ("Figure_5_pk_dual_screen.png", "Fig. 5 | P. knowlesi dual-screen condition comparison."),
        ("Figure_6_species_condition_contrast.png", "Fig. 6 | Descriptive species-versus-condition discordance contrast."),
    ]
    docx_from_markdownish(OUT / "manuscript" / "Nature_Microbiology_replacement_manuscript.docx", "Species- and condition-aware essentiality mapping across malaria parasites", body, figs)

    si = f"""
# Supplementary Information

## Supplementary Note 1: Freeze and audit
Release identifier: `v20261017_5fold_9route_lightgbm`. The main matrix has 628 raw features. The split is the original 95%-deduplicated label-stratified random split. The freeze manifest records matrix, identifier and output hashes.

## Supplementary Note 2: Historical assessment boundary
The current run does not use test labels for fitting, thresholding or candidate selection. However, the route profile was selected after historical inspection of Pk results. Main test scores are therefore post-hoc frozen evaluation, not untouched prospective validation.

## Supplementary Note 3: Pk condition analysis
The primary Pk definite-call flip rate is 12.57% (440/3,501; 95% CI 11.51–13.71%). Spearman correlation of continuous scores is 0.744 (95% bootstrap CI 0.730–0.757). Exact A1H sequence-level prediction is not claimed.

## Supplementary Note 4: Four-species LOSO
The LOSO table includes Pf, Pb, Pk and Tg. Tg is retained as a full held-out target and is not silently excluded because it is phylogenetically more distant.

## Supplementary Note 5: Interpretation
Directional groups, target overlays and website calls are computational hypotheses. They require experimental follow-up and should not be interpreted as direct evidence of drug efficacy or parasite-selective safety.
"""
    write_text(OUT / "supplementary_information" / "Supplementary_Information.md", si)
    docx_from_markdownish(OUT / "supplementary_information" / "Supplementary_Information.docx", "Supplementary Information", si, [])

    cover = f"""
# Cover letter

Dear Editors,

We submit **Species- and condition-aware essentiality mapping across malaria parasites** for consideration as a Research Article in *Nature Microbiology*. The study addresses a central problem in parasite functional genomics: experimental essentiality labels are informative across species, but their transfer is limited by species, host-cell condition and assay design.

The manuscript contributes three linked analyses. First, a frozen five-fold nine-route LightGBM model uses the original 628-feature matrix and reproduces held-out ranking across P. falciparum, P. berghei and P. knowlesi. Second, a four-species leave-one-species-out analysis includes T. gondii as an explicit distant held-out target. Third, two independent P. knowlesi screens under different red-cell conditions quantify within-species disagreement and temper interpretation of cross-species directional groups.

The release is accompanied by hash-indexed source data, reproducibility code, six figures, a model card and a versioned ProtoEss Atlas. We state the historical non-blind evaluation limitation and do not present computational calls as experimental essentiality evidence.

This work fits the journal's interest in protozoan biology, host interactions, omics, computational biology and evolution. The manuscript is not under consideration elsewhere. All authors have approved the submission and declare no competing interests.

Sincerely,\nYuan-Nong Ye and Zhu Zeng\nfor all authors
"""
    write_text(OUT / "submission" / "Cover_Letter_Nature_Microbiology.md", cover)

    reporting = """
# Reporting Summary

## Study design
- Study type: computational reanalysis and predictive modelling of public parasite essentiality screens.
- Organisms: P. falciparum, P. berghei, P. knowlesi and T. gondii for LOSO; two P. knowlesi screen conditions for the condition analysis.
- No new animal or human experiments.

## Data and analysis
- Main split: 95% redundancy reduction followed by label-stratified random train/validation/test partition.
- Main model: five-fold, nine-route LightGBM using 628 raw features.
- Dynamic label-derived features: fold-local reconstruction from fitting labels only.
- Independent analysis: four-species LOSO including Tg.
- Condition analysis: retrospective ortholog-mapped Pk dual-screen comparison; not exact A1H sequence-level prediction.

## Reproducibility
- Release manifest: `manifest/final_freeze_manifest.json`.
- All published tables and figures have SHA-256 entries.
- Test labels are excluded from current fitting, thresholding and candidate selection; historical route selection was not fully blind.
"""
    write_text(OUT / "submission" / "Reporting_Summary.md", reporting)

    checklist = """
# Machine-learning checklist

- [x] Task and labels documented.
- [x] Original 95%-deduplicated, label-stratified random split retained.
- [x] Raw 628 features retained; no PCA or manual homology30/50 removal.
- [x] Five-fold OOF structure and random seed recorded.
- [x] Dynamic label-derived features rebuilt fold-locally.
- [x] Thresholds derived from training-only probabilities.
- [x] Test labels excluded from current fit, threshold and candidate selection.
- [x] Historical test exposure disclosed.
- [x] LOSO includes Pf, Pb, Pk and Tg.
- [x] Pk human-RBC analysis separated from main evaluation.
- [x] Exact A1H sequence-level prediction not claimed without complete A1H features.
- [x] Model-derived calls labelled computational prioritisation.
- [x] Source data, code, figures and hashes packaged.
"""
    write_text(OUT / "submission" / "Machine_Learning_Checklist.md", checklist)

    mapping = pd.DataFrame([
        ("main model", "historical five-branch/older route", "five-fold nine-route LightGBM", "main results and figures"),
        ("features", "mixed historical matrices", "raw 628-feature repaired matrix", "main results"),
        ("split", "historical iterations", "original 95%-deduplicated label-stratified random split", "main results"),
        ("LOSO", "malaria-only and older clade variants", "four-species table including Tg", "companion analysis"),
        ("Pk human-RBC", "historical label-assisted route", "separate retrospective ortholog-mapped condition analysis", "supplementary / limitation"),
        ("directional wording", "lineage-specific", "species- and condition-associated", "main text"),
    ], columns=["component", "old_version", "release_version", "placement"])
    mapping.to_csv(OUT / "old_vs_new" / "old_vs_new_mapping.tsv", sep="\t", index=False)

    print(f"wrote manuscript and submission documents to {OUT}")


if __name__ == "__main__":
    main()
